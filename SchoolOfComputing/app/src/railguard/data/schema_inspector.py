"""Bounded-memory CSV/XLSX schema inspection."""

from __future__ import annotations

import csv
import re
import warnings
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from railguard.utils.hashing import stable_hash
from railguard.data.timestamp import DOOR_TIMESTAMP, parse_door_timestamps

TIME_HINT = re.compile(r"time|date|timestamp", re.I)
TARGET_HINT = re.compile(r"label|target|damage|status|fault", re.I)


def _csv_shape(path: Path) -> tuple[int, int]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader, [])
        rows = sum(1 for _ in reader)
    return rows, len(header)


def _summarize_frame(frame: pd.DataFrame) -> dict[str, Any]:
    numeric = frame.select_dtypes(include=np.number)
    stats: dict[str, Any] = {}
    for column in numeric.columns:
        series = numeric[column].replace([np.inf, -np.inf], np.nan)
        stats[str(column)] = {
            "min": float(series.min()) if series.notna().any() else None,
            "max": float(series.max()) if series.notna().any() else None,
            "mean": float(series.mean()) if series.notna().any() else None,
            "std": float(series.std()) if series.notna().sum() > 1 else None,
        }
    columns = [str(c) for c in frame.columns]
    return {
        "columns": columns,
        "dtypes": {str(k): str(v) for k, v in frame.dtypes.items()},
        "null_percent": {str(k): float(v) for k, v in (frame.isna().mean() * 100).items()},
        "constant_columns_in_sample": [str(c) for c in frame if frame[c].nunique(dropna=False) <= 1],
        "numeric_columns": [str(c) for c in numeric.columns],
        "non_numeric_columns": [str(c) for c in frame.columns if c not in numeric.columns],
        "candidate_timestamp_columns": [c for c in columns if TIME_HINT.search(c)],
        "candidate_target_leakage_columns": [c for c in columns if TARGET_HINT.search(c)],
        "numeric_summary": stats,
        "schema_hash": stable_hash({"columns": columns, "dtypes": {str(k): str(v) for k, v in frame.dtypes.items()}}),
    }


def _timestamp_diagnostics(frame: pd.DataFrame, candidates: list[str]) -> dict[str, Any]:
    diagnostics: dict[str, Any] = {}
    for column in candidates:
        strings = frame[column].astype(str)
        if len(strings) and strings.map(lambda value: bool(DOOR_TIMESTAMP.fullmatch(value))).mean() > 0.9:
            values, errors = parse_door_timestamps(strings, strict=False)
            parsed = pd.Series(pd.to_datetime(values), index=frame.index)
        else:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                parsed = pd.to_datetime(frame[column], errors="coerce")
        valid = parsed.dropna()
        if len(valid) < 2:
            continue
        delta = valid.diff().dropna().dt.total_seconds()
        diagnostics[column] = {
            "parse_success": float(parsed.notna().mean()),
            "duplicated": int(valid.duplicated().sum()),
            "monotonic": bool(valid.is_monotonic_increasing),
            "sampling_seconds": {"median": float(delta.median()), "min": float(delta.min()), "max": float(delta.max())},
        }
    return diagnostics


def inspect_csv(path: str | Path, sample_rows: int = 2000) -> dict[str, Any]:
    source = Path(path)
    n_rows, n_columns = _csv_shape(source)
    frame = pd.read_csv(source, nrows=sample_rows)
    report = _summarize_frame(frame)
    report.update({"path": str(source), "extension": source.suffix.lower(), "n_rows": n_rows, "n_columns": n_columns, "sampled_rows": len(frame)})
    report["timestamp_diagnostics"] = _timestamp_diagnostics(frame, report["candidate_timestamp_columns"])
    return report


def inspect_xlsx(path: str | Path, sample_rows: int = 2000) -> dict[str, Any]:
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        values = sheet.iter_rows(values_only=True)
        header = [str(v) if v is not None else f"unnamed_{i}" for i, v in enumerate(next(values, []))]
        rows = [row for _, row in zip(range(sample_rows), values, strict=False)]
        frame = pd.DataFrame(rows, columns=header)
        summary = _summarize_frame(frame)
        summary.update({"name": sheet.title, "n_rows": max(sheet.max_row - 1, 0), "n_columns": sheet.max_column, "sampled_rows": len(frame)})
        summary["timestamp_diagnostics"] = _timestamp_diagnostics(frame, summary["candidate_timestamp_columns"])
        sheets.append(summary)
    workbook.close()
    return {"path": str(source), "extension": source.suffix.lower(), "sheets": sheets, "sheet_names": [s["name"] for s in sheets]}


def inspect_file(path: str | Path, sample_rows: int = 2000) -> dict[str, Any]:
    source = Path(path)
    if source.suffix.lower() in {".xlsx", ".xlsm"}:
        return inspect_xlsx(source, sample_rows)
    if source.suffix.lower() in {".csv", ".txt"}:
        return inspect_csv(source, sample_rows)
    raise ValueError(f"Unsupported inspection format: {source.suffix}")


def reports_to_markdown(reports: list[dict[str, Any]]) -> str:
    lines = ["# Schema inspection report", ""]
    extensions = Counter(r["extension"] for r in reports)
    lines.append(f"Files inspected: {len(reports)} ({dict(extensions)})")
    for report in reports:
        lines.extend(["", f"## `{report['path']}`"])
        if "sheets" in report:
            lines.append(f"Sheets: {', '.join(report['sheet_names'])}")
            for sheet in report["sheets"]:
                lines.append(f"- {sheet['name']}: {sheet['n_rows']} rows × {sheet['n_columns']} columns")
        else:
            lines.append(f"Shape: {report['n_rows']} rows × {report['n_columns']} columns")
            lines.append(f"Columns: {', '.join(report['columns'])}")
    return "\n".join(lines) + "\n"
